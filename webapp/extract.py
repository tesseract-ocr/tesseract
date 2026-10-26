"""Turn an uploaded file (image, PDF, DOCX, XLSX) into text, using OCR where needed."""
import io
import os
import re
import subprocess
import tempfile

import docx
import fitz  # PyMuPDF
import openpyxl

TESSERACT = os.environ.get("TESSERACT_BIN", "tesseract")
TESSDATA = os.environ.get("TESSDATA_PREFIX")

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif", ".webp"}
SUPPORTED = IMAGE_EXT | {".pdf", ".docx", ".xlsx"}
LEGACY = {".doc", ".xls"}
MIN_TEXT_LAYER = 20  # PDF pages with fewer characters than this are treated as scans


# Tesseract puts a space between CJK characters; drop spaces that sit between two of them.
_CJK = r"\u3000-\u303f\u3040-\u30ff\u3400-\u4dbf\u4e00-\u9fff\uff00-\uffef"
_CJK_GAP = re.compile(rf"(?<=[{_CJK}])[ \t]+(?=[{_CJK}])")


def clean_cjk(text):
    return _CJK_GAP.sub("", text)


def ocr_path(path, lang):
    cmd = [TESSERACT, path, "-", "-l", lang]
    if TESSDATA:
        cmd += ["--tessdata-dir", TESSDATA]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
    if r.returncode != 0:
        raise RuntimeError(r.stderr.strip() or "tesseract failed")
    return clean_cjk(r.stdout.strip())


def ocr_bytes(data, lang, suffix=".png"):
    with tempfile.NamedTemporaryFile(suffix=suffix) as tmp:
        tmp.write(data)
        tmp.flush()
        return ocr_path(tmp.name, lang)


def _pdf(path, lang):
    parts = []
    with fitz.open(path) as doc:
        for i, page in enumerate(doc, 1):
            text = page.get_text().strip()
            if len(text) < MIN_TEXT_LAYER:
                text = ocr_bytes(page.get_pixmap(dpi=300).tobytes("png"), lang)
            parts.append(f"--- Page {i} ---\n{text}")
    return "\n\n".join(parts)


def _docx(path, lang):
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append("\t".join(c.text.strip() for c in row.cells))
    for rel in d.part.rels.values():
        if "image" in rel.reltype:
            blob = rel.target_part.blob
            ext = os.path.splitext(rel.target_part.partname)[1] or ".png"
            if ext.lower() in IMAGE_EXT:
                txt = ocr_bytes(blob, lang, ext)
                if txt:
                    parts.append(f"[image text]\n{txt}")
    return "\n".join(parts)


def _xlsx(path, lang):
    wb = openpyxl.load_workbook(path, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"--- Sheet: {ws.title} ---")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                parts.append("\t".join("" if v is None else str(v) for v in row))
        for img in getattr(ws, "_images", []):
            txt = ocr_bytes(img._data(), lang)
            if txt:
                parts.append(f"[image text]\n{txt}")
    return "\n".join(parts)


def extract(path, filename, lang):
    ext = os.path.splitext(filename)[1].lower()
    if ext in LEGACY:
        raise ValueError(f"{ext} is a legacy format; save it as {ext}x (or PDF) and upload again")
    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported file type '{ext}'. Use: " + ", ".join(sorted(SUPPORTED)))
    if ext in IMAGE_EXT:
        return ocr_path(path, lang)
    return {".pdf": _pdf, ".docx": _docx, ".xlsx": _xlsx}[ext](path, lang)
